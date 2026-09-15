#!/usr/bin/env python3
"""Build the Whoniverse spreadsheet from the ledger TSVs.

The TSVs in this directory are the source of truth. This script validates them
and writes out/whoniverse.xlsx, which is what gets imported into the public
Google Sheet, plus out/file-names.tsv for the addon work.

    python ledger/build.py

Nothing here talks to Google. Import the xlsx by hand:
File, Import, Upload, then "Replace spreadsheet".
"""
import json, os, re, sys, unicodedata
from collections import Counter, OrderedDict

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit('openpyxl is needed: pip install openpyxl')

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(HERE, 'out')

CATEGORY_ORDER = ['Main Show', 'Special', 'Minisode', 'Animated Series',
                  'Prequel', 'Animated Restoration', 'Movie']
COLOUR = {
    'Main Show': 'b6d7a8', 'Special': 'f4cccc', 'Minisode': 'a4c2f4',
    'Animated Series': 'f9cb9c', 'Prequel': 'ffe599',
    'Animated Restoration': 'f9cb9c', 'Movie': 'b6d7a8',
}
DESCRIPTION = {
    'Main Show': 'Main episodes of the show.',
    'Special': 'Full episodes between seasons and such.',
    'Minisode': 'Mini episodes that are story relevant.',
    'Animated Series': 'Short animated series.',
    'Prequel': 'Prequel episodes of other main story episodes.',
    'Animated Restoration': 'Missing episodes surviving as BBC animations.',
    'Movie': 'The 1996 television film.',
    'Best Found': 'The best copy worth getting: its source, resolution and audio, '
                  'and the day the indexers were last asked. Not the largest file that '
                  'exists anywhere — everything here is re-encoded for streaming, so a '
                  'disc remux is not a target.',
    'We Have': 'What the file we hold actually is. Blank means nothing downloaded yet.',
}
HEADER_FILL = 'd9d9d9'
SEASON_EVEN, SEASON_ODD = 'd9d9d9', 'b7b7b7'
UNAVAILABLE = 'b4a7d6'
FILE_FILL = 'efefef'
STATUS_FILL = {
    'ok': 'e2efda',        # matches the best that exists
    'below': 'fff2cc',     # a better release exists
    'cadence': 'fce4d6',   # converted from another standard
    'upscale': 'ddebf7',   # bigger than the real source, so an upscale
    'missing': 'fce4e4',   # nothing downloaded yet
}
ALL_WHO_NAME = 'Complete Chronology'


def read_tsv(path):
    # Every row must have exactly as many columns as the header. This used to
    # pad short rows with empty strings, which hid a lost tab: The Christmas
    # Invasion ran `checked` and `have` together into one cell, so best stayed
    # right, checked became "2026-09-151872x1080 23.976fps AAC", released took
    # the description and the description went empty, and the build said
    # nothing. A dropped tab is a typo, not a shape to accommodate.
    with open(path, encoding='utf8') as fh:
        rows = [ln.rstrip('\n').split('\t') for ln in fh if ln.strip('\n')]
    head, body = rows[0], rows[1:]
    ragged = ['%s line %d: %d columns, header has %d'
              % (os.path.basename(path), i, len(r), len(head))
              for i, r in enumerate(body, 2) if len(r) != len(head)]
    if ragged:
        raise SystemExit('ragged TSV, a tab is missing or doubled:\n  '
                         + '\n  '.join(ragged))
    return [OrderedDict(zip(head, r)) for r in body]


def load():
    # series.tsv is the registry: one row per series, holding both how the
    # sheet renders it and everything the addon needs to describe it. The two
    # used to live apart, the sheet here and the prose in lib/series.js, which
    # is how the film ended up with no description at all.
    tabs = []
    for t in read_tsv(os.path.join(HERE, 'series.tsv')):
        tabs.append(dict(key=t['key'], name=t['sheet_name'],
                         numbered=t['numbered'] == '1',
                         cats=t['categories'].split('|'),
                         folder=t.get('folder', ''),
                         stremio_id=t.get('stremio_id', ''),
                         data=t.get('data', ''),
                         art=t.get('art', ''),
                         release_info=t.get('release_info', ''),
                         genres=[g for g in t.get('genres', '').split('|') if g],
                         description=t.get('description', ''),
                         # The chronology has no series file: it is the whole
                         # of all-who.tsv, so it is validated and numbered
                         # from there rather than from a file of its own.
                         derived=t.get('derived', '') == '1'))
    series = {}
    for fn in sorted(os.listdir(os.path.join(HERE, 'series'))):
        if fn.endswith('.tsv'):
            series[fn.split('-', 1)[1][:-4]] = read_tsv(os.path.join(HERE, 'series', fn))
    allwho = read_tsv(os.path.join(HERE, 'all-who.tsv'))
    return tabs, series, allwho


def validate(tabs, series, allwho):
    errs = []
    keys = [t['key'] for t in tabs if not t['derived']]
    for k in series:
        if k not in keys:
            errs.append('series file %s has no row in series.tsv' % k)
    for t in tabs:
        if t['derived']:
            continue
        rows = series.get(t['key'])
        if rows is None:
            errs.append('series.tsv lists %s but there is no series file' % t['key']); continue
        seen = Counter()
        for i, r in enumerate(rows, 2):
            where = '%s line %d' % (t['key'], i)
            if r['category'] not in t['cats']:
                errs.append('%s: category %r is not one of %s' % (where, r['category'], t['cats']))
            if r['category'] not in CATEGORY_ORDER:
                errs.append('%s: unknown category %r' % (where, r['category']))
            if r['status'] not in ('ok', 'missing'):
                errs.append('%s: status must be ok or missing, got %r' % (where, r['status']))
            if (r['status'] == 'missing') != bool(r['note'].strip()):
                errs.append('%s: a missing row needs a note, an ok row must not have one' % where)
            if r['season'] != '-' and not r['season'].isdigit():
                errs.append('%s: season must be a number or -' % where)
            if not r['title'].strip():
                errs.append('%s: empty title' % where)
            seen[(r['season'], r['category'], r['title'])] += 1
        for (s, cat, ti), n in seen.items():
            if n > 1:
                errs.append('%s: %r (%s) appears %d times in season %s' % (t['key'], ti, cat, n, s))
    have = Counter((r['series'], r['season'], r['category'], r['title']) for r in allwho)
    want = Counter((t['key'], r['season'], r['category'], r['title'])
                   for t in tabs if not t['derived']
                   for r in series.get(t['key'], []))
    for k in want - have:
        errs.append('all-who.tsv is missing %s' % (k,))
    for k in have - want:
        errs.append('all-who.tsv has %s, which is in no series file' % (k,))
    return errs


def number(tabs, series):
    """episode numbers per season, in file order"""
    for t in tabs:
        if t['derived']:
            continue
        counter = Counter()
        for r in series[t['key']]:
            if r['season'] == '-':
                r['ep'] = None
            else:
                counter[r['season']] += 1
                r['ep'] = counter[r['season']]


NO_SUFFIX = ('Main Show', 'Movie')


def cell_text(r):
    t = r['title']
    if r['category'] not in NO_SUFFIX:
        t += ' (%s)' % r['category']
    if r['status'] == 'missing':
        t += ' ⚠️ (%s)' % r['note']
    return t


def slug(text):
    text = re.sub(r'⚠️?\s*\([^)]*\)$', '', text).strip()
    text = unicodedata.normalize('NFKD', text).encode('ascii', 'ignore').decode().lower()
    text = text.replace(' ', '_')
    return re.sub(r'_+$', '', re.sub(r'[^a-z0-9_]', '', text))


BAD_CADENCE = (23.976, 29.97, 30.3, 23.98, 29.976)


def _height(text):
    """The vertical resolution a string is talking about, as a 16:9 equivalent.

    A 2:1 frame like 1920x960 is a 1080-class picture with the letterbox baked
    out, not a shortfall, so it is compared on the height a 16:9 frame of the
    same width would have. Frames narrower than 16:9, such as a 720x576 PAL
    DVD, keep their real height.
    """
    if not text:
        return None
    m = re.search(r'(\d{3,4})\s*[pi]\b', text)
    if m:
        return int(m.group(1))
    m = re.search(r'(\d{3,4})x(\d{3,4})', text)
    if m:
        w, h = int(m.group(1)), int(m.group(2))
        return max(h, round(w * 9 / 16))
    if re.search(r'\bSD\b', text):
        return 576
    return None


def quality(r, series_key=''):
    """Compare what we hold against the best copy worth getting.

    Returns (status, reason). Status is one of missing, below, cadence,
    upscale, ok. Only ok means there is nothing left to do.

    There used to be two targets here — what exists anywhere, and what a search
    had actually turned up — and rows were held against both. That produced 523
    permanently red rows chasing disc remuxes that would never ship: every file
    in the bucket is re-encoded for streaming, so a 90GB season pack is not a
    goal, it is a different hobby. One target now, `best`, describing the source
    and resolution worth having.
    """
    have = (r.get('have') or '')
    if not have:
        return ('missing', 'nothing downloaded')
    target = (r.get('best') or '')
    want, got = _height(target), _height(have)
    fps = re.search(r'([\d.]+)fps', have)
    fps = float(fps.group(1)) if fps else None
    if want and got and got < want:
        return ('below', 'have %dp, %s is worth getting' % (got, target))
    if fps and round(fps, 3) in BAD_CADENCE and series_key != 'wilderness-years':
        # A converted cadence is only a fault while a clean copy can still be
        # had. Once a search has established that what we hold IS the best
        # obtainable, the cadence is a fact about the release rather than a job
        # on the list, and colouring hundreds of rows amber for it buries the
        # rows that can actually be fixed.
        if want and got and got >= want:
            return ('ok', 'the best worth getting, though %gfps is a conversion' % fps)
        return ('cadence', '%gfps, converted from another standard' % fps)
    if want and got and got > want:
        return ('upscale', 'have %dp, above the %s we record' % (got, target))
    return ('ok', 'matches the best worth getting')


def hunt(r):
    """Has anyone actually checked this row, and was the answer any good?

    Returns (worth_searching, reason). `best` is only ever what somebody knew on
    the day it was written, so the question this answers is not "does something
    better exist" — something better always exists — but "is this row's target
    stale". A row nobody has asked the indexers about is the one to ask about.
    """
    on = (r.get('checked') or '').strip()
    best = (r.get('best') or '').strip()
    if not on:
        return (True, 'never searched')
    if not best:
        return (True, 'searched %s, nothing to aim at recorded' % on)
    return (False, '')


def safe_title(name):
    """xlsx bans a handful of punctuation in sheet names and caps them at 31
    characters. Google's own export does the same, so a tab may need renaming
    after import."""
    clean = re.sub(r'[:\\/?*\[\]]', '', name).strip()
    return clean[:31].strip()


def layout(tab):
    """column numbers: season, ep, {category: col}, file, then the two quality
    columns in the order they are read: the best copy worth getting, and what we
    hold."""
    if not tab['numbered']:
        return dict(season=None, ep=None, cats={tab['cats'][0]: 2}, file=4,
                    best=6, have=8, width=8)
    cats = {c: 6 + 2 * i for i, c in enumerate(tab['cats'])}
    last = 6 + 2 * (len(tab['cats']) - 1)
    if tab['name'] == ALL_WHO_NAME:
        return dict(season=2, ep=4, cats=cats, file=None, best=None,
                    have=None, width=last)
    return dict(season=2, ep=4, cats=cats, file=last + 2,
                best=last + 4, have=last + 6,
                width=last + 6)


def formula(tab, lay, last_row):
    if lay['file'] is None:
        return None
    if not tab['numbered']:
        return ('=ARRAYFORMULA(IF(B5<>"", REGEXREPLACE(REGEXREPLACE(SUBSTITUTE('
                'LOWER(B5), " ", "_"), "[^a-z0-9_]", ""), "_+$", ""), ""))')
    cats = '&'.join('{0}5:{0}{1}'.format(get_column_letter(c), last_row)
                    for c in lay['cats'].values())
    return ('=ARRAYFORMULA(IF((B5:B{L}<>"")*(D5:D{L}<>""), "S"&TEXT(B5:B{L}, "00")'
            '&"_E"&TEXT(D5:D{L}, "00")&"_"&REGEXREPLACE(REGEXREPLACE(SUBSTITUTE('
            'LOWER(REGEXREPLACE({C}, "⚠️\\s*\\([^)]*\\)$", "")), " ", "_"), '
            '"[^a-z0-9_]", ""), "_+$", ""), ""))').format(L=last_row, C=cats)


def paint(ws, tab, rows, lay):
    body = Font(name='Inter', size=10)
    head = Font(name='Inter', size=10, bold=True)
    align = Alignment(horizontal='left', vertical='center', wrap_text=False)
    grey = PatternFill('solid', fgColor=HEADER_FILL)

    for col in range(1, lay['width'] + 1):
        letter = get_column_letter(col)
        if not tab['numbered']:
            ws.column_dimensions[letter].width = 1.13 if col % 2 else (40.13 if col == 2 else 34.0)
        else:
            wide = 50.13
            if col in (2, 4):
                wide = 6.13
            elif col in (lay.get('best'), lay.get('have')):
                wide = 34.0
            ws.column_dimensions[letter].width = 1.13 if col % 2 else wide
    ws.row_dimensions[1].height = 6
    ws.row_dimensions[2].height = 36
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 6
    ws.freeze_panes = 'A4'

    headers = []
    if tab['numbered']:
        headers += [(lay['season'], 'Sn.', None), (lay['ep'], 'Ep.', None)]
    headers += [(c, name, DESCRIPTION[name]) for name, c in lay['cats'].items()]
    if lay['file']:
        headers.append((lay['file'], 'File Name', None))
    if lay.get('best'):
        headers.append((lay['best'], 'Best Found', DESCRIPTION['Best Found']))
    if lay.get('have'):
        headers.append((lay['have'], 'We Have', DESCRIPTION['We Have']))
    for col, name, desc in headers:
        for row, value in ((2, name), (3, desc)):
            cell = ws.cell(row, col)
            cell.value = value
            cell.font = head if row == 2 else body
            cell.alignment = align
            cell.fill = grey

    for i, r in enumerate(rows):
        row = 5 + i
        ws.row_dimensions[row].height = 18
        gone = r['status'] == 'missing'
        last_cat = max(lay['cats'].values())
        if gone:
            for col in range(2, last_cat + 1):
                ws.cell(row, col).fill = PatternFill('solid', fgColor=UNAVAILABLE)
        if tab['numbered'] and r['ep'] is not None:
            shade = SEASON_EVEN if int(r['sn']) % 2 == 0 else SEASON_ODD
            for col, value in ((lay['season'], int(r['sn'])), (lay['ep'], r['ep'])):
                cell = ws.cell(row, col)
                cell.value = value
                cell.fill = PatternFill('solid', fgColor=UNAVAILABLE if gone else shade)
        cell = ws.cell(row, lay['cats'][r['category']])
        cell.value = cell_text(r)
        cell.fill = PatternFill('solid', fgColor=UNAVAILABLE if gone else COLOUR[r['category']])
        if lay['file']:
            f = ws.cell(row, lay['file'])
            f.fill = PatternFill('solid', fgColor=FILE_FILL)
            if i == 0:
                f.value = formula(tab, lay, 4 + len(rows))
        if lay.get('best'):
            c = ws.cell(row, lay['best'])
            on = r.get('checked') or ''
            c.value = ('%s  (%s)' % (r.get('best'), on) if r.get('best') and on
                       else (r.get('best') or None))
            # Amber where nobody has asked the indexers yet, which is a gap in
            # the work rather than a fault in the file.
            worth, _why = hunt(r)
            c.fill = PatternFill('solid', fgColor=STATUS_FILL['below'] if worth else FILE_FILL)
        if lay.get('have'):
            status, why = quality(r, tab['key'])
            c = ws.cell(row, lay['have'])
            c.value = r.get('have') or None
            c.fill = PatternFill('solid', fgColor=STATUS_FILL[status])
        for col in range(1, lay['width'] + 1):
            ws.cell(row, col).font = body
            ws.cell(row, col).alignment = align


def registry(tabs):
    """Everything the addon needs to describe itself, straight from the ledger.

    Written into data/ rather than out/ because the addon loads it at runtime
    and out/ is generated output that the deploy ignores.
    """
    addon = {r['field']: r['value'] for r in read_tsv(os.path.join(HERE, 'addon.tsv'))}
    missing = [t['key'] for t in tabs if not t['description']]
    if missing:
        raise SystemExit('these series have no description in series.tsv: %s'
                         % ', '.join(missing))
    out = {
        'addon': addon,
        'series': [
            OrderedDict((
                ('key', t['key']), ('stremioId', t['stremio_id']),
                ('name', t['name'].replace(' ✅', '').strip()),
                ('folder', t['folder']), ('data', t['data']), ('art', t['art']),
                ('releaseInfo', t['release_info']), ('genres', t['genres']),
                ('description', t['description']),
            )) for t in tabs
        ],
    }
    path = os.path.join(ROOT, 'data', 'registry.json')
    with open(path, 'w', encoding='utf8', newline='\n') as fh:
        json.dump(out, fh, indent=1, ensure_ascii=False)
        fh.write('\n')


def main():
    tabs, series, allwho = load()
    errs = validate(tabs, series, allwho)
    if errs:
        print('%d problem(s) in the ledger:' % len(errs))
        for e in errs[:40]:
            print('  ' + e)
        sys.exit(1)
    number(tabs, series)

    # The chronology has no rows of its own: it is every other series in
    # broadcast order, so it is skipped everywhere a series file is read.
    real = [t for t in tabs if not t['derived']]
    by_tab = {t['key']: t for t in real}
    lookup = {}
    for t in real:
        for r in series[t['key']]:
            lookup[(t['key'], r['season'], r['category'], r['title'])] = r

    # All Who: running season number per (series, season), first appearance wins
    running, order = {}, []
    all_rows = []
    for a in allwho:
        src = lookup[(a['series'], a['season'], a['category'], a['title'])]
        key = (a['series'], a['season'])
        if a['season'] != '-' and key not in running:
            running[key] = len(running) + 1
            order.append((running[key], by_tab[a['series']]['name'], a['season']))
        row = dict(src)
        row['sn'] = running.get(key) if a['season'] != '-' else None
        row['ep'] = src['ep']
        all_rows.append(row)

    wb = Workbook()
    wb.remove(wb.active)
    counts, renamed = [], []
    for t in real:
        rows = [dict(r, sn=r['season']) for r in series[t['key']]]
        ws = wb.create_sheet(safe_title(t['name']))
        if safe_title(t['name']) != t['name']:
            renamed.append((t['name'], safe_title(t['name'])))
        paint(ws, t, rows, layout(t))
        counts.append((t['name'], len(rows)))
    aw = dict(key='all-who', name=ALL_WHO_NAME, numbered=True, cats=CATEGORY_ORDER)
    ws = wb.create_sheet(ALL_WHO_NAME)
    paint(ws, aw, all_rows, layout(aw))
    counts.append((ALL_WHO_NAME, len(all_rows)))

    os.makedirs(OUT, exist_ok=True)
    xlsx = os.path.join(OUT, 'whoniverse.xlsx')
    wb.save(xlsx)

    # released and description ride along so anything building a data file has
    # the whole row and does not have to go back to the series TSV for the two
    # columns the addon actually shows a viewer.
    names = ['series\tseason\tepisode\tcategory\tstatus\ttitle\tfile_name\t'
             'best	checked	have	released	description	note']
    for t in real:
        for r in series[t['key']]:
            fn = ''
            if r['ep'] is not None:
                fn = (slug(cell_text(r)) if not t['numbered']
                      else 'S%02d_E%02d_%s' % (int(r['season']), r['ep'], slug(cell_text(r))))
            names.append('\t'.join([t['key'], r['season'], str(r['ep'] or ''),
                                    r['category'], r['status'], r['title'], fn,
                                    r.get('best') or '', r.get('checked') or '',
                                    r.get('have') or '',
                                    r.get('released') or '',
                                    (r.get('description') or '').replace('\t', ' '),
                                    (r.get('note') or '').replace('\t', ' ')]))
    with open(os.path.join(OUT, 'file-names.tsv'), 'w', encoding='utf8', newline='\n') as fh:
        fh.write('\n'.join(names) + '\n')

    for name, n in counts:
        print('  %-40s %5d rows' % (name, n))
    if renamed:
        print('\n  xlsx cannot hold these tab names, rename them after importing:')
        for was, now in renamed:
            print('    %-42s -> %s' % (was, now))
    print('\n  %d series-season blocks, %d rows total' % (len(running), len(all_rows)))
    registry(tabs)
    print('  wrote %s' % xlsx)
    print('  wrote %s' % os.path.join(OUT, 'file-names.tsv'))
    print('  wrote %s' % os.path.join(ROOT, 'data', 'registry.json'))


if __name__ == '__main__':
    main()
